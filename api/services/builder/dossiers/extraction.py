"""[.mark] Builder, lot 3 : le texte d'un document déposé, page par page, extrait SUR PLACE (rien n'est envoyé à personne).

Formats : PDF (``pypdf``, BSD-3), Word ``.docx`` (``python-docx``, MIT), Markdown et texte. Chaque page garde son
numéro : c'est la source d'une ligne de la fiche (« audit.pdf, p. 3 »). Un Word et un texte n'ont pas de pages :
ils sont coupés en parties d'une taille fixe, numérotées (« section 2 »).

Un document **illisible** (chiffré, scanné sans texte, corrompu, format inconnu, trop gros) est SIGNALÉ avec sa raison
et ne rend aucune page : on n'invente jamais ce qu'on n'a pas lu.
"""

from __future__ import annotations

import asyncio
import io
import os
import zipfile
from dataclasses import dataclass, field

from loguru import logger

OCTETS_MAX = 15 * 1024 * 1024
PAGES_MAX = 300
CARACTERES_MAX = 400_000
PARTIE_CARACTERES = 3000
DECOMPRESSE_MAX = 80 * 1024 * 1024  # un .docx est une archive : on refuse une bombe de décompression

EXTENSIONS_TEXTE = (".txt", ".md", ".markdown")
FORMATS = (".pdf", ".docx", *EXTENSIONS_TEXTE)


@dataclass
class Page:
    numero: int
    texte: str
    etiquette: str = "p."  # « p. 3 » pour un PDF, « section 2 » pour le reste


@dataclass
class Extraction:
    pages: list[Page] = field(default_factory=list)
    illisible: str | None = None  # la raison, en clair, quand rien n'a pu être lu
    tronque: bool = False

    @property
    def caracteres(self) -> int:
        return sum(len(p.texte) for p in self.pages)


def _illisible(raison: str) -> Extraction:
    return Extraction(illisible=raison)


def _en_parties(texte: str) -> list[Page]:
    """Un texte sans pages, coupé aux fins de paragraphe en parties numérotées."""
    blocs = [b for b in texte.replace("\r\n", "\n").split("\n\n") if b.strip()]
    parties: list[str] = []
    courante = ""
    for bloc in blocs:
        while len(bloc) > PARTIE_CARACTERES:  # un paragraphe géant est coupé net : mieux que le perdre
            if courante:
                parties.append(courante)
                courante = ""
            parties.append(bloc[:PARTIE_CARACTERES])
            bloc = bloc[PARTIE_CARACTERES:]
        if courante and len(courante) + len(bloc) + 2 > PARTIE_CARACTERES:
            parties.append(courante)
            courante = ""
        courante = f"{courante}\n\n{bloc}" if courante else bloc
    if courante:
        parties.append(courante)
    return [Page(numero=i, texte=t.strip(), etiquette="section") for i, t in enumerate(parties, start=1) if t.strip()]


def _decoder(contenu: bytes) -> str:
    for codage in ("utf-8-sig", "cp1252"):
        try:
            return contenu.decode(codage)
        except UnicodeDecodeError:
            continue
    return contenu.decode("utf-8", errors="replace")


def _pdf(contenu: bytes) -> Extraction:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        lecteur = PdfReader(io.BytesIO(contenu))
        if lecteur.is_encrypted:
            return _illisible("The PDF is protected by a password: remove the protection and deposit it again.")
        pages: list[Page] = []
        tronque = False
        for numero, page in enumerate(lecteur.pages, start=1):
            if numero > PAGES_MAX:
                tronque = True
                break
            texte = (page.extract_text() or "").strip()
            if texte:
                pages.append(Page(numero=numero, texte=texte))
    except (PyPdfError, ValueError, KeyError, OSError, RecursionError, TypeError, AttributeError) as erreur:
        logger.info(f"[builder] PDF illisible : {type(erreur).__name__}")
        return _illisible("The PDF could not be read (damaged or unsupported file).")
    if not pages:
        return _illisible("The PDF has no text layer (a scan?): deposit a version with selectable text.")
    return Extraction(pages=pages, tronque=tronque)


def _word(contenu: bytes) -> Extraction:
    try:
        with zipfile.ZipFile(io.BytesIO(contenu)) as archive:
            if sum(info.file_size for info in archive.infolist()) > DECOMPRESSE_MAX:
                return _illisible("The Word file is too large once unpacked.")
        from docx import Document

        document = Document(io.BytesIO(contenu))
        morceaux = [p.text for p in document.paragraphs if p.text.strip()]
        for tableau in document.tables:
            for ligne in tableau.rows:
                cellules = [c.text.strip() for c in ligne.cells if c.text.strip()]
                if cellules:
                    morceaux.append(" | ".join(cellules))
    except Exception as erreur:  # noqa: BLE001 -- python-docx lève des erreurs de plusieurs familles sur un fichier abîmé
        logger.info(f"[builder] Word illisible : {type(erreur).__name__}")
        return _illisible("The Word file could not be read (damaged, or not a .docx).")
    return _depuis_texte("\n\n".join(morceaux), "The Word file holds no text.")


def _depuis_texte(texte: str, si_vide: str) -> Extraction:
    texte = texte.strip()
    if not texte:
        return _illisible(si_vide)
    return Extraction(pages=_en_parties(texte))


def borner(resultat: Extraction) -> Extraction:
    """Plafonne le texte total d'un document (la fiche et le cerveau ont une taille limite)."""
    if resultat.illisible or resultat.caracteres <= CARACTERES_MAX:
        return resultat
    gardees: list[Page] = []
    total = 0
    for page in resultat.pages:
        if total + len(page.texte) > CARACTERES_MAX:
            break
        gardees.append(page)
        total += len(page.texte)
    return Extraction(pages=gardees, tronque=True)


def extraire(nom: str, contenu: bytes) -> Extraction:
    """Le texte de ``contenu`` selon l'extension de ``nom``. Ne lève jamais : un fichier illisible est un résultat."""
    if len(contenu) > OCTETS_MAX:
        return _illisible(f"The file is too large ({len(contenu) // (1024 * 1024)} MB, at most {OCTETS_MAX // (1024 * 1024)} MB).")
    if not contenu:
        return _illisible("The file is empty.")
    extension = os.path.splitext(nom or "")[1].lower()
    if extension == ".pdf":
        resultat = _pdf(contenu)
    elif extension == ".docx":
        resultat = _word(contenu)
    elif extension in EXTENSIONS_TEXTE:
        resultat = _depuis_texte(_decoder(contenu), "The file holds no text.")
    else:
        return _illisible(f"Format not supported ({extension or 'no extension'}): PDF, Word (.docx), Markdown or text.")
    return borner(resultat)


DELAI_EXTRACTION_S = 60.0


async def extraire_hors_boucle(nom: str, contenu: bytes, delai_s: float | None = None) -> Extraction:
    """``extraire`` dans un fil à part, avec un délai (revue R3) : une conversion PDF ou Word est du calcul pur, qui gelait la boucle
    de l'API (et les appels du même processus) tant qu'elle durait. Au-delà du délai, le fichier est dit illisible en le disant ;
    ⚠️ le fil lui-même ne peut pas être tué et finit seul, mais plus rien ne l'attend. Ne lève jamais."""
    delai = DELAI_EXTRACTION_S if delai_s is None else delai_s
    try:
        return await asyncio.wait_for(asyncio.to_thread(extraire, nom, contenu), timeout=delai)
    except TimeoutError:
        return _illisible(f"Reading the file took more than {delai:g} seconds: split it or send a simpler file.")


def source(nom: str, page: Page) -> str:
    """La source telle qu'elle s'écrit sur une ligne de la fiche : « audit.pdf, p. 3 » ou « notes.md, section 2 »."""
    if page.etiquette == "p.":
        return f"{nom}, p. {page.numero}"
    return f"{nom}, {page.etiquette} {page.numero}"
