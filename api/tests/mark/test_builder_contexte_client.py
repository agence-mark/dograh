"""[.mark] builder-agent-vocal, lot 3 : les dossiers déposés, la fiche sourcée, les faits classés, la base de connaissances.

Ce qui doit tenir :
* chaque format (PDF, Word, Markdown, texte) est lu sur place, page par page ; un document illisible est SIGNALÉ avec sa
  raison, jamais inventé ;
* un site se lit en pages publiques seulement : aucune adresse interne (localhost, réseaux privés, métadonnées, rebond
  DNS, redirection vers l'interne, schémas autres que http/https, identifiants, ports inhabituels), bornes de taille et
  de pages, robots.txt respecté ;
* dès qu'il y a des dépôts, une ligne de la fiche cite un dépôt (ou la réponse d'une personne) ; un trou reste ouvert tant
  qu'aucune ligne sourcée ne le comble et la fiche ne se valide pas avec un trou ;
* un fait est classé objet, base de connaissances ou règle ; un fait en base n'est pas un objet ; l'objet créé par le
  builder est rendu dans le prompt par ``{{nom}}`` ; base du client injoignable = ``{{nom}}`` inchangé, l'appel continue ;
* la base de connaissances est proposée avec des extraits mot pour mot, jamais un chiffre ou une promesse absents de
  l'audit ; rien n'est rattaché sans la validation d'une personne.
"""

from __future__ import annotations

import io
from datetime import datetime
from types import SimpleNamespace

import httpx
import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient

from api.routes import builder as route
from api.routes import builder_contexte as route_contexte
from api.routes import objets_agent as route_objets
from api.schemas.objets_agent import ObjetAgent, ObjetEnvoye, ObjetsAgent
from api.services.auth.depends import get_user, get_user_with_selected_organization
from api.services.builder import brouillon, conception, fiche
from api.services.builder.actions.registre import Contexte, executer
from api.services.builder.dossiers import base, depot, extraction, site
from api.services.objets_agent.appel import injecter_objets
from api.tests.mark.builder_commun import AGENT, ORG_AUTRE, ORG_CLIENT, utilisateur
from api.utils.template_renderer import render_template

EQUIPE = Contexte(user=utilisateur(1), workflow_id=AGENT)


# --- un stockage de fichiers en mémoire -------------------------------------------------------------------------------


@pytest.fixture(autouse=True)
def _stockage(monde, monkeypatch):
    fichiers: dict[str, bytes] = {}

    async def ecrire(cle, donnees):
        fichiers[cle] = donnees

    async def lire(cle):
        return fichiers.get(cle)

    async def supprimer(cle):
        fichiers.pop(cle, None)

    monkeypatch.setattr(depot, "ecrire_octets", ecrire)
    monkeypatch.setattr(depot, "lire_octets", lire)
    monkeypatch.setattr(depot, "supprimer_octets", supprimer)
    monkeypatch.setattr(route_contexte, "db_client", monde)
    monkeypatch.setattr(route, "db_client", monde)
    return fichiers


# --- des fichiers d'essai -----------------------------------------------------------------------------------------------


def un_pdf(pages: list[str]) -> bytes:
    """Un vrai PDF à une police standard, une page par texte (pypdf relit la table des objets)."""
    objets: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    nombre = len(pages)
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(nombre))
    objets.append(f"<< /Type /Pages /Kids [{kids}] /Count {nombre} >>".encode())
    police = 3 + 2 * nombre
    for i, texte in enumerate(pages):
        flux = f"BT /F1 12 Tf 72 720 Td ({texte}) Tj ET".encode()
        objets.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R /Resources << /Font << /F1 {police} 0 R >> >> >>".encode()
        )
        objets.append(b"<< /Length %d >>\nstream\n" % len(flux) + flux + b"\nendstream")
    objets.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    sortie = io.BytesIO(b"%PDF-1.4\n")
    positions = []
    for i, corps in enumerate(objets, start=1):
        positions.append(sortie.tell())
        sortie.write(f"{i} 0 obj\n".encode() + corps + b"\nendobj\n")
    debut = sortie.tell()
    sortie.write(f"xref\n0 {len(objets) + 1}\n0000000000 65535 f \n".encode())
    for position in positions:
        sortie.write(f"{position:010d} 00000 n \n".encode())
    sortie.write(f"trailer\n<< /Size {len(objets) + 1} /Root 1 0 R >>\nstartxref\n{debut}\n%%EOF".encode())
    return sortie.getvalue()


def un_word(paragraphes: list[str]) -> bytes:
    from docx import Document

    document = Document()
    for p in paragraphes:
        document.add_paragraph(p)
    tableau = document.add_table(rows=1, cols=2)
    tableau.rows[0].cells[0].text = "Jour"
    tableau.rows[0].cells[1].text = "Lundi au vendredi"
    sortie = io.BytesIO()
    document.save(sortie)
    return sortie.getvalue()


AUDIT = "Nous ouvrons du lundi au vendredi de 9h a 18h. Le devis est realise sur place."
AUDIT_P2 = "La zone couvre le departement. Les marques suivies sont Alpha et Beta."


# --- 1. l'extraction ----------------------------------------------------------------------------------------------------


def test_un_pdf_est_lu_page_par_page_avec_sa_source():
    resultat = extraction.extraire("audit.pdf", un_pdf([AUDIT, AUDIT_P2]))
    assert resultat.illisible is None and [p.numero for p in resultat.pages] == [1, 2]
    assert "lundi au vendredi" in resultat.pages[0].texte and "Alpha" in resultat.pages[1].texte
    assert extraction.source("audit.pdf", resultat.pages[1]) == "audit.pdf, p. 2"


def test_un_word_est_lu_paragraphes_et_tableaux_en_sections():
    resultat = extraction.extraire("notes.docx", un_word(["Premier paragraphe de l'audit.", "Second paragraphe."]))
    texte = " ".join(p.texte for p in resultat.pages)
    assert "Premier paragraphe" in texte and "Lundi au vendredi" in texte  # le tableau est lu aussi
    assert extraction.source("notes.docx", resultat.pages[0]) == "notes.docx, section 1"


@pytest.mark.parametrize("nom", ["a.md", "a.txt", "a.markdown"])
def test_markdown_et_texte_sont_lus(nom):
    resultat = extraction.extraire(nom, "# Horaires\n\nOuvert de 9h à 18h.\n".encode())
    assert resultat.illisible is None and "9h à 18h" in resultat.pages[0].texte


def test_un_texte_ancien_en_cp1252_n_est_pas_perdu():
    assert "é" in extraction.extraire("a.txt", "Réception à 9h".encode("cp1252")).pages[0].texte


def test_un_long_texte_est_coupe_en_parties_numerotees():
    long = "\n\n".join(f"Paragraphe numéro {i} " + "x" * 400 for i in range(30))
    pages = extraction.extraire("a.md", long.encode()).pages
    assert len(pages) > 3 and [p.numero for p in pages] == list(range(1, len(pages) + 1))


@pytest.mark.parametrize(
    ("nom", "contenu", "morceau"),
    [
        ("casse.pdf", b"%PDF-1.4 ce n'est pas un pdf", "could not be read"),
        ("scan.pdf", None, "no text layer"),
        ("faux.docx", b"PK pas une archive", "could not be read"),
        ("vide.md", b"   \n  ", "no text"),
        ("rien.txt", b"", "empty"),
        ("tableur.xlsx", b"PK", "not supported"),
        ("enorme.txt", b"x" * (extraction.OCTETS_MAX + 1), "too large"),
    ],
)
def test_un_document_illisible_est_signale_avec_sa_raison_et_ne_rend_aucune_page(nom, contenu, morceau):
    if contenu is None:  # un PDF d'une page blanche : pas de couche de texte
        from pypdf import PdfWriter

        ecrivain = PdfWriter()
        ecrivain.add_blank_page(200, 200)
        tampon = io.BytesIO()
        ecrivain.write(tampon)
        contenu = tampon.getvalue()
    resultat = extraction.extraire(nom, contenu)
    assert resultat.pages == [] and morceau in (resultat.illisible or "")


def test_un_pdf_protege_par_mot_de_passe_est_signale():
    from pypdf import PdfReader, PdfWriter

    ecrivain = PdfWriter(clone_from=PdfReader(io.BytesIO(un_pdf([AUDIT]))))
    try:
        ecrivain.encrypt("secret", algorithm="RC4-128")
    except Exception:  # noqa: BLE001 -- sans la bibliothèque de chiffrement, ce cas ne se fabrique pas ici
        pytest.skip("encryption not available in this environment")
    tampon = io.BytesIO()
    ecrivain.write(tampon)
    resultat = extraction.extraire("protege.pdf", tampon.getvalue())
    assert resultat.pages == [] and "password" in (resultat.illisible or "")


def test_une_archive_word_qui_se_decompresse_demesurement_est_refusee(monkeypatch):
    monkeypatch.setattr(extraction, "DECOMPRESSE_MAX", 100)
    assert "too large once unpacked" in (extraction.extraire("gros.docx", un_word(["a " * 500])).illisible or "")


# --- 2. le dépôt ---------------------------------------------------------------------------------------------------------


async def test_un_fichier_depose_est_lu_garde_et_retrouve_sa_source(monde, _stockage):
    depose = await depot.deposer_fichier(ORG_CLIENT, AGENT, "audit.pdf", un_pdf([AUDIT, AUDIT_P2]), "Marie")
    assert depose["statut"] == "lu" and depose["pages"] == 2 and "cle_texte" not in depose and "cle_fichier" not in depose
    assert any(cle.endswith("/original") for cle in _stockage)  # l'original est gardé
    (enregistre,) = await depot.lister(ORG_CLIENT, AGENT)
    pages = await depot.pages_de(ORG_CLIENT, enregistre)
    assert "Alpha" in pages[1].texte and depot.etiquette_de_source(enregistre, pages[1]) == "audit.pdf, p. 2"


async def test_un_document_illisible_reste_dans_la_liste_avec_sa_raison(monde):
    depose = await depot.deposer_fichier(ORG_CLIENT, AGENT, "scan.pdf", b"%PDF-1.4 abime", "Marie")
    assert depose["statut"] == "illisible" and depose["raison"] and depose["pages"] == 0
    assert await depot.pages_de(ORG_CLIENT, (await depot.lister(ORG_CLIENT, AGENT))[0]) == []


async def test_le_nom_est_nettoye_les_doublons_refuses_et_le_retrait_efface(monde, _stockage):
    depose = await depot.deposer_fichier(ORG_CLIENT, AGENT, "..\\..\\etc/passwd\x00.md", b"un texte", "Marie")
    assert depose["nom"] == "passwd.md"
    with pytest.raises(depot.ErreurDepot, match="already deposited"):
        await depot.deposer_fichier(ORG_CLIENT, AGENT, "PASSWD.md", b"autre", "Marie")
    assert await depot.retirer(ORG_CLIENT, AGENT, depose["id"]) and not _stockage
    assert await depot.lister(ORG_CLIENT, AGENT) == [] and not await depot.retirer(ORG_CLIENT, AGENT, "inconnu")


async def test_les_depots_d_un_agent_sont_bornes(monde, monkeypatch):
    monkeypatch.setattr(depot, "DEPOTS_MAX", 2)
    await depot.deposer_fichier(ORG_CLIENT, AGENT, "a.md", b"a", "M")
    await depot.deposer_fichier(ORG_CLIENT, AGENT, "b.md", b"b", "M")
    with pytest.raises(depot.ErreurDepot, match="At most 2"):
        await depot.deposer_fichier(ORG_CLIENT, AGENT, "c.md", b"c", "M")


# --- 3. le site : pages publiques seulement ----------------------------------------------------------------------------------

PUBLIC = "93.184.216.34"


async def dns_public(hote):
    return [PUBLIC]


def page(corps: str, type_="text/html; charset=utf-8", statut=200, **entetes):
    return httpx.Response(statut, headers={"content-type": type_, **entetes}, content=corps.encode())


def faux_site(pages: dict[str, httpx.Response], vus: list | None = None):
    def repondre(requete: httpx.Request) -> httpx.Response:
        if vus is not None:
            vus.append(requete)
        reponse = pages.get(requete.url.path)
        return reponse if reponse is not None else httpx.Response(404)

    return httpx.MockTransport(repondre)


ACCUEIL = (
    "<html><head><title>Accueil</title><style>.x{}</style><script>alert(1)</script></head>"
    "<body><h1>Bienvenue</h1><p>Ouvert du lundi au vendredi.</p>"
    '<a href="/contact">Contact</a><a href="https://autre.example/page">Externe</a><a href="mailto:a@b.c">mail</a></body></html>'
)


async def test_le_html_est_reduit_a_son_texte_sans_scripts_ni_styles():
    titre, texte, liens = site.lire_le_html(ACCUEIL)
    assert titre == "Accueil" and "Bienvenue" in texte and "lundi au vendredi" in texte
    assert "alert" not in texte and ".x{}" not in texte and "/contact" in liens


async def test_un_site_public_est_lu_en_suivant_ses_liens_du_meme_site_seulement():
    vus: list[httpx.Request] = []
    transport = faux_site({"/": page(ACCUEIL), "/contact": page("<p>Écrire à l'équipe.</p>"), "/robots.txt": httpx.Response(404)}, vus)
    lecture = await site.lire_le_site("https://client.example/", resolveur=dns_public, transport=transport)
    assert [p.adresse for p in lecture.pages] == ["https://client.example/", "https://client.example/contact"]
    assert not any("autre.example" in str(r.url) or r.headers["host"] == "autre.example" for r in vus)
    # La connexion va à l'adresse VÉRIFIÉE (pas de seconde résolution), le nom reste dans l'en-tête Host.
    assert all(r.url.host == PUBLIC and r.headers["host"] == "client.example" for r in vus)


@pytest.mark.parametrize(
    "adresse",
    [
        "ftp://client.example/",
        "file:///etc/passwd",
        "javascript:alert(1)",
        "https://user:mdp@client.example/",
        "https://client.example:8080/",
        "https://client.example:22/",
        "http://localhost/",
        "http://127.0.0.1/",
        "http://[::1]/",
        "http://10.0.0.5/",
        "http://192.168.1.10/",
        "http://172.16.0.1/",
        "http://169.254.169.254/latest/meta-data/",
        "http://0.0.0.0/",
        "http://[::ffff:127.0.0.1]/",
        "http://[fd00::1]/",
        "http://100.64.0.1/",
    ],
)
async def test_une_adresse_interne_ou_un_schema_etranger_est_refuse(adresse):
    appels = []
    transport = faux_site({}, appels)

    async def resolveur_reel_simule(hote):
        return {"localhost": ["127.0.0.1"]}.get(hote, [PUBLIC])

    with pytest.raises(site.SiteRefuse):
        await site.lire_le_site(adresse, resolveur=resolveur_reel_simule, transport=transport)
    assert appels == []  # rien n'est parti sur le réseau


async def test_un_nom_qui_se_resout_vers_le_reseau_interne_est_refuse_meme_s_il_a_l_air_public():
    async def rebond(hote):
        return ["93.184.216.34", "10.0.0.7"]  # une seule adresse privée parmi les réponses suffit à refuser

    appels = []
    with pytest.raises(site.SiteRefuse, match="private or internal"):
        await site.lire_le_site("https://piege.example/", resolveur=rebond, transport=faux_site({}, appels))
    assert appels == []


async def test_une_redirection_vers_l_interne_est_refusee_a_l_etape():
    async def resolveur(hote):
        return ["10.0.0.1"] if hote == "interne.example" else [PUBLIC]

    transport = faux_site({"/": page("", statut=302, location="https://interne.example/admin"), "/robots.txt": httpx.Response(404)})
    with pytest.raises(site.SiteRefuse, match="private or internal"):
        await site.lire_le_site("https://client.example/", resolveur=resolveur, transport=transport)


async def test_les_redirections_sont_bornees():
    transport = faux_site({"/": page("", statut=302, location="/"), "/robots.txt": httpx.Response(404)})
    with pytest.raises(site.SiteRefuse, match="Too many redirections"):
        await site.lire_le_site("https://client.example/", resolveur=dns_public, transport=transport)


async def test_une_page_trop_grosse_ou_non_textuelle_est_refusee():
    gros = page("<p>" + "x" * (site.OCTETS_PAGE_MAX + 10) + "</p>")
    with pytest.raises(site.SiteRefuse, match="too large"):
        await site.lire_le_site("https://client.example/", resolveur=dns_public, transport=faux_site({"/": gros}))
    pdf = httpx.Response(200, headers={"content-type": "application/pdf"}, content=b"%PDF")
    with pytest.raises(site.SiteRefuse, match="Not a text page"):
        await site.lire_le_site("https://client.example/", resolveur=dns_public, transport=faux_site({"/": pdf}))


async def test_le_nombre_de_pages_est_borne():
    liens = "".join(f'<a href="/p{i}">p</a>' for i in range(30))
    pages = {"/": page(f"<p>Accueil</p>{liens}")} | {f"/p{i}": page(f"<p>Page {i}</p>") for i in range(30)}
    lecture = await site.lire_le_site("https://client.example/", resolveur=dns_public, transport=faux_site(pages), pages_max=4)
    assert len(lecture.pages) == 4


async def test_robots_txt_est_respecte():
    robots = httpx.Response(200, headers={"content-type": "text/plain"}, content=b"User-agent: *\nDisallow: /prive\n")
    pages = {"/": page('<p>Accueil</p><a href="/prive">x</a><a href="/public">y</a>'), "/prive": page("<p>Secret</p>"), "/public": page("<p>Ouvert</p>"), "/robots.txt": robots}
    lecture = await site.lire_le_site("https://client.example/", resolveur=dns_public, transport=faux_site(pages))
    assert [p.adresse.rsplit("/", 1)[-1] for p in lecture.pages] == ["", "public"]
    assert any("robots.txt" in raison for _, raison in lecture.refusees)
    tout_interdit = httpx.Response(200, headers={"content-type": "text/plain"}, content=b"User-agent: *\nDisallow: /\n")
    with pytest.raises(site.SiteRefuse, match="robots.txt"):
        await site.lire_le_site("https://client.example/", resolveur=dns_public, transport=faux_site({"/robots.txt": tout_interdit}))


async def test_un_site_depose_devient_un_depot_dont_chaque_page_a_son_adresse(monde):
    transport = faux_site({"/": page(ACCUEIL), "/contact": page("<p>Écrire à l'équipe.</p>"), "/robots.txt": httpx.Response(404)})
    depose = await depot.deposer_site(ORG_CLIENT, AGENT, "client.example", "Marie", resolveur=dns_public, transport=transport)
    assert depose["statut"] == "lu" and depose["genre"] == "site" and depose["nom"] == "https://client.example"
    (enregistre,) = await depot.lister(ORG_CLIENT, AGENT)
    pages = await depot.pages_de(ORG_CLIENT, enregistre)
    assert depot.etiquette_de_source(enregistre, pages[1]) == "https://client.example/contact"
    with pytest.raises(depot.ErreurDepot, match="private or internal"):
        await depot.deposer_site(ORG_CLIENT, AGENT, "http://127.0.0.1/", "Marie")
    assert len(await depot.lister(ORG_CLIENT, AGENT)) == 1  # le refus n'a rien rangé


# --- 4. la fiche sourcée -----------------------------------------------------------------------------------------------------


async def deposer_l_audit():
    return await depot.deposer_fichier(ORG_CLIENT, AGENT, "audit.pdf", un_pdf([AUDIT, AUDIT_P2]), "Marie")


async def test_tirer_des_documents_rend_les_pages_avec_leur_source_et_signale_l_illisible(monde):
    await deposer_l_audit()
    await depot.deposer_fichier(ORG_CLIENT, AGENT, "scan.pdf", b"%PDF-1.4 abime", "Marie")
    resultat = await executer("fiche_tirer_des_documents", {}, EQUIPE)
    assert resultat["ok"] and [e["source"] for e in resultat["extraits"]] == ["audit.pdf, p. 1", "audit.pdf, p. 2"]
    assert resultat["illisibles"][0]["nom"] == "scan.pdf" and resultat["illisibles"][0]["raison"]
    assert "horaires" in resultat["champs_standard_sans_ligne"]  # ce qui manque devient une question
    assert resultat["carte"]["type"] == "depot"


async def test_la_lecture_est_bornee_et_se_continue(monde, monkeypatch):
    from api.services.builder.actions import contexte_client

    monkeypatch.setattr(contexte_client, "BUDGET_CARACTERES", 100)
    await depot.deposer_fichier(ORG_CLIENT, AGENT, "audit.pdf", un_pdf([AUDIT, AUDIT_P2, "Troisieme page de l'audit."]), "M")
    premier = await executer("fiche_tirer_des_documents", {}, EQUIPE)
    assert len(premier["extraits"]) < 3 and premier["suite"]
    suivant = await executer("fiche_tirer_des_documents", {"a_partir_de": premier["suite"]["a_partir_de"]}, EQUIPE)
    assert suivant["extraits"][0]["source"] != premier["extraits"][0]["source"]


async def test_un_depot_inconnu_est_refuse(monde):
    resultat = await executer("fiche_tirer_des_documents", {"documents": ["nope"]}, EQUIPE)
    assert resultat["ok"] is False and "Unknown deposit" in resultat["erreur"]


async def test_une_ligne_doit_citer_un_depot_une_fois_les_dossiers_deposes(monde):
    await deposer_l_audit()
    inventee = await executer("fiche_ecrire", {"champs": {"horaires": {"valeur": "9h-18h", "source": "ma memoire"}}}, EQUIPE)
    assert inventee["ok"] is False and inventee["code"] == "source_inconnue"
    bonne = await executer("fiche_ecrire", {"champs": {"horaires": {"valeur": "9h-18h", "source": "audit.pdf, p. 1"}}}, EQUIPE)
    assert bonne["ok"] is True
    personne = await executer("fiche_ecrire", {"champs": {"zone": {"valeur": "le département", "source": "Réponse de Marie"}}}, EQUIPE)
    assert personne["ok"] is True


async def test_un_trou_reste_ouvert_et_empeche_la_validation_puis_la_personne_le_ferme(monde):
    await deposer_l_audit()
    await executer(
        "fiche_ecrire",
        {
            "champs": {"horaires": {"valeur": "9h-18h", "source": "audit.pdf, p. 1"}},
            "trous": [{"champ": "marques", "question": "Quelles marques ?", "reponse_suggeree": "Alpha, Beta"}],
        },
        EQUIPE,
    )
    c = client_pour(utilisateur(1))
    assert c.post(f"/builder/fiche/{AGENT}/valider").status_code == 422  # un trou ouvert : rien n'est validé
    vue = c.put(f"/builder/fiche/{AGENT}/ligne", json={"champ": "marques", "valeur": "Alpha, Beta"}).json()
    assert vue["fiche"]["trous"] == [] and vue["fiche"]["champs"]["marques"]["source"].startswith("Réponse de ")
    assert c.post(f"/builder/fiche/{AGENT}/valider").status_code == 200
    # Corriger une ligne retire la validation : ce qui est validé n'est jamais modifié en silence.
    c.put(f"/builder/fiche/{AGENT}/ligne", json={"champ": "horaires", "valeur": "8h-17h"})
    assert (await fiche.lire(ORG_CLIENT, AGENT)).validee is False


# --- 5. les faits classés : objet, base de connaissances ou règle ----------------------------------------------------------------


async def test_chaque_fait_a_une_place_celle_du_champ_ou_celle_de_sa_taille(monde):
    resultat = await executer(
        "fiche_ecrire",
        {
            "champs": {
                "horaires": {"valeur": "9h-18h", "source": "s"},
                "motifs_d_appel": {"valeur": "devis, urgence", "source": "s"},
                "garantie_pieces": {"valeur": "x" * 50, "source": "s"},
                "catalogue": {"valeur": "y" * 900, "source": "s"},
                "tarif_special": {"valeur": "50 euros", "source": "s", "classement": "base"},
            }
        },
        EQUIPE,
    )
    assert resultat["carte"]["classements"] == {
        "horaires": "objet",
        "motifs_d_appel": "regle",
        "garantie_pieces": "objet",
        "catalogue": "base",
        "tarif_special": "base",
    }
    assert (await executer("fiche_lire", {}, EQUIPE))["classements"]["catalogue"] == "base"


async def test_un_fait_trop_gros_n_est_pas_un_objet(monde):
    gros = await executer("fiche_ecrire", {"champs": {"catalogue": {"valeur": "z" * 3000, "source": "s", "classement": "objet"}}}, EQUIPE)
    assert gros["ok"] is False and "knowledge base" in gros["erreur"]


class FauxObjets:
    """Les objets d'une base de client en mémoire, derrière les mêmes fonctions que les routes (mêmes garde-fous)."""

    RESERVES = {"date", "adresse"}

    def __init__(self):
        self.objets: dict[int, ObjetAgent] = {}

    async def get_objets(self, user):
        return ObjetsAgent(objets=list(self.objets.values()), plafond_caracteres=3000)

    async def _ecrire(self, user, objet: ObjetEnvoye, objet_id):
        if objet.nom in self.RESERVES:
            raise HTTPException(status_code=409, detail=f"« {objet.nom} » is a variable the call sets itself")
        if objet.forme == "valeur" and objet.type is None:
            raise HTTPException(status_code=422, detail="A single value has a type")
        identifiant = objet_id or len(self.objets) + 1
        ecrit = ObjetAgent(**objet.model_dump(), id=identifiant, modifie_le=datetime.now(), taille=len(objet.valeur or ""))
        self.objets[identifiant] = ecrit
        return ecrit


@pytest.fixture
def objets(monkeypatch):
    faux = FauxObjets()
    monkeypatch.setattr(route_objets, "get_objets", faux.get_objets)
    monkeypatch.setattr(route_objets, "_ecrire", faux._ecrire)
    return faux


HORAIRES = {"nom": "horaires", "libelle": "Horaires", "forme": "valeur", "type": "texte", "valeur": "du lundi au vendredi, 9h-18h"}


async def test_un_objet_cree_par_le_builder_est_rendu_dans_le_prompt(monde, objets):
    resultat = await executer("objets_agent_ecrire", {"objets": [HORAIRES]}, EQUIPE)
    assert resultat["ok"] and resultat["objets"][0]["action"] == "cree" and resultat["objets"][0]["placeholder"] == "{{horaires}}"
    # Réécrire le même nom modifie, ne duplique pas.
    resultat = await executer("objets_agent_ecrire", {"objets": [{**HORAIRES, "valeur": "9h-17h"}]}, EQUIPE)
    assert resultat["objets"][0]["action"] == "modifie" and len(objets.objets) == 1
    listes = await executer("objets_agent_lister", {}, EQUIPE)
    assert [o["nom"] for o in listes["objets"]] == ["horaires"]
    # Rendu : l'injection existante du fork pose l'objet, le prompt le cite par {{nom}}.
    lecture = SimpleNamespace(objets_etat="lu", objets=list(objets.objets.values()), objets_noms=["horaires"])
    contexte, _ = injecter_objets({}, lecture)
    assert render_template("Ce que tu sais : {{horaires}}", contexte) == "Ce que tu sais : 9h-17h"


async def test_les_garde_fous_des_routes_s_appliquent_aux_objets_du_builder(monde, objets):
    reserve = await executer("objets_agent_ecrire", {"objets": [{**HORAIRES, "nom": "date"}, {**HORAIRES, "nom": "zone"}]}, EQUIPE)
    assert reserve["ok"] is False
    par_nom = {r["nom"]: r for r in reserve["objets"]}
    assert par_nom["date"]["ok"] is False and par_nom["date"]["code"] == 409 and par_nom["zone"]["ok"] is True  # un refus n'arrête pas les autres
    sans_type = await executer("objets_agent_ecrire", {"objets": [{"nom": "marques", "libelle": "M", "forme": "valeur"}]}, EQUIPE)
    assert sans_type["ok"] is False
    # Pas d'objet à déclencheur ni chargé à la demande : le schéma refuse tout champ inconnu.
    declencheur = await executer("objets_agent_ecrire", {"objets": [{**HORAIRES, "declencheur": "si on demande"}]}, EQUIPE)
    assert declencheur["ok"] is False and "Invalid arguments" in declencheur["erreur"]
    mauvais_nom = await executer("objets_agent_ecrire", {"objets": [{**HORAIRES, "nom": "Horaires Du Jour"}]}, EQUIPE)
    assert mauvais_nom["ok"] is False


async def test_sans_base_rattachee_le_builder_le_dit_au_lieu_de_faire_semblant(monde, monkeypatch):
    async def sans_base(user):
        raise HTTPException(status_code=409, detail="Attach a client database first: the objects live there.")

    monkeypatch.setattr(route_objets, "get_objets", sans_base)
    resultat = await executer("objets_agent_lister", {}, EQUIPE)
    assert resultat["ok"] is False and "client database" in resultat["erreur"]


def test_base_du_client_injoignable_le_nom_reste_ecrit_tel_quel_et_l_appel_continue():
    lecture = SimpleNamespace(objets_etat="injoignable", objets=[], objets_noms=["horaires"])
    contexte, estampille = injecter_objets({}, lecture)
    assert render_template("Ce que tu sais : {{horaires}}", contexte) == "Ce que tu sais : {{horaires}}"
    assert estampille["etat"] == "injoignable"


async def test_un_modele_de_la_bibliotheque_s_applique_avec_les_valeurs_du_client(monde, objets):
    from api.services.builder.bibliotheque import service

    modele = {"objets": [{"nom": "zone", "libelle": "Zone", "forme": "valeur", "type": "texte", "valeur": ""},
                         {"nom": "marques", "libelle": "Marques", "forme": "liste", "colonnes": [{"libelle": "Marque", "type": "texte"}]}]}
    await service.ecrire_candidat("objets-agent-types", "zone-et-marques", "Zone et marques", modele, auteur="Evan")
    resultat = await executer(
        "objets_agent_ecrire",
        {"depuis": "zone-et-marques", "valeurs": {"zone": "le département"}, "lignes": {"marques": [{"Marque": "Alpha"}, {"Marque": "Beta"}]}},
        EQUIPE,
    )
    assert resultat["ok"], resultat
    ecrits = {o.nom: o for o in objets.objets.values()}
    assert ecrits["zone"].valeur == "le département" and [l.valeurs["Marque"] for l in ecrits["marques"].lignes] == ["Alpha", "Beta"]
    inconnu = await executer("objets_agent_ecrire", {"depuis": "n-existe-pas"}, EQUIPE)
    assert inconnu["ok"] is False


async def test_les_faits_classes_objet_deviennent_des_objets_et_pas_ceux_de_la_base(monde, objets):
    await executer(
        "fiche_ecrire",
        {
            "champs": {
                "horaires": {"valeur": "9h-18h", "source": "s"},
                "marques": {"valeur": [{"Marque": "Alpha", "Zone": "Nord"}, {"Marque": "Beta", "Zone": "Sud"}], "source": "s"},
                "catalogue": {"valeur": "y" * 900, "source": "s", "classement": "base"},
                "motifs_d_appel": {"valeur": "devis", "source": "s"},
            }
        },
        EQUIPE,
    )
    avant = await executer("fiche_creer_les_objets", {}, EQUIPE)
    assert avant["ok"] is False and avant["code"] == "fiche_non_validee" and objets.objets == {}
    await fiche.valider(ORG_CLIENT, AGENT, "Marie")
    apres = await executer("fiche_creer_les_objets", {}, EQUIPE)
    assert apres["ok"] and sorted(o["nom"] for o in apres["objets"]) == ["horaires", "marques"]
    assert apres["ailleurs"] == {"base": ["catalogue"], "regle": ["motifs_d_appel"]}  # un fait, une place
    assert {o.nom for o in objets.objets.values()} == {"horaires", "marques"}
    liste = next(o for o in objets.objets.values() if o.nom == "marques")
    assert liste.forme == "liste" and [c.libelle for c in liste.colonnes] == ["Marque", "Zone"]


# --- 6. la base de connaissances ----------------------------------------------------------------------------------------------

FAQ_DEPOT = (
    "Questions fréquentes. Le devis est réalisé sur place et il est gratuit. "
    "Nous intervenons sous 48 heures dans tout le département. Le paiement se fait à la livraison."
)


async def deposer_la_faq():
    await depot.deposer_fichier(ORG_CLIENT, AGENT, "faq.md", FAQ_DEPOT.encode(), "Marie")


def proposition(**surcharges):
    return {
        "titre": "Devis et délais",
        "contenu": "Le devis est réalisé sur place et il est gratuit. Nous intervenons sous 48 heures.",
        "sources": [{"document": "faq.md", "extrait": "Le devis est réalisé sur place et il est gratuit. Nous intervenons sous 48 heures"}],
        **surcharges,
    }


async def test_un_document_de_base_est_propose_avec_ses_extraits_mot_pour_mot(monde):
    await deposer_la_faq()
    resultat = await executer("base_proposer", proposition(), EQUIPE)
    assert resultat["ok"] and resultat["etat"] == "propose" and resultat["carte"]["type"] == "base"
    (document,) = await base.lister(ORG_CLIENT, AGENT)
    assert document.etat == "propose" and document.document_uuid is None  # rien n'est créé dans la base de connaissances


@pytest.mark.parametrize(
    ("surcharge", "morceau"),
    [
        ({"contenu": "Le devis coûte 120 euros et nous intervenons sous 48 heures."}, "Number(s) 120"),
        ({"contenu": "Le devis est gratuit, nous intervenons sous 48 heures, satisfait ou remboursé."}, "Commitment word"),
        ({"contenu": "Le devis est offert et garanti."}, "Commitment word"),
        ({"sources": [{"document": "faq.md", "extrait": "Le devis est réalisé chez vous en une heure chrono"}]}, "not found word for word"),
        ({"sources": [{"document": "inconnu.pdf", "extrait": "Le devis est réalisé sur place"}]}, "not a deposited document"),
        ({"sources": [{"document": "faq.md", "page": 9, "extrait": "Le devis est réalisé sur place"}]}, "no page 9"),
    ],
)
async def test_jamais_un_tarif_un_delai_ou_une_promesse_absents_de_l_audit(monde, surcharge, morceau):
    await deposer_la_faq()
    resultat = await executer("base_proposer", proposition(**surcharge), EQUIPE)
    assert resultat["ok"] is False and morceau in resultat["erreur"]
    assert await base.lister(ORG_CLIENT, AGENT) == []  # rien n'est rangé


async def test_les_chiffres_de_l_audit_passent_quel_que_soit_leur_ecriture(monde):
    await depot.deposer_fichier(ORG_CLIENT, AGENT, "tarifs.md", "Le forfait coûte 1 200 euros, livré en 3 jours.".encode(), "M")
    ok = await executer(
        "base_proposer",
        {"titre": "Forfait", "contenu": "Forfait à 1200 euros, livré en 3 jours.", "sources": [{"document": "tarifs.md", "extrait": "Le forfait coûte 1 200 euros, livré en 3 jours."}]},
        EQUIPE,
    )
    assert ok["ok"], ok


async def test_une_proposition_remplace_la_precedente_du_meme_titre_tant_qu_elle_n_est_pas_validee(monde):
    await deposer_la_faq()
    premiere = await executer("base_proposer", proposition(), EQUIPE)
    seconde = await executer("base_proposer", proposition(contenu="Le devis est réalisé sur place."), EQUIPE)
    assert premiere["id"] == seconde["id"] and len(await base.lister(ORG_CLIENT, AGENT)) == 1


# --- les routes : les gestes humains ----------------------------------------------------------------------------------------


def client_pour(user) -> TestClient:
    app = FastAPI()
    app.include_router(route.router)
    app.dependency_overrides[get_user_with_selected_organization] = lambda: user
    app.dependency_overrides[get_user] = lambda: user
    return TestClient(app)


@pytest.fixture
def document_cree(monkeypatch):
    class Crees(list):
        modes: list[str]

    crees = Crees()
    modes = []

    async def creer(user, titre, contenu, mode="chunked"):
        crees.append((user.selected_organization_id, titre, contenu))
        modes.append(mode)
        return f"uuid-{len(crees)}"

    monkeypatch.setattr(base, "creer_le_document_de_base", creer)
    crees.modes = modes
    return crees


async def approuver_la_conception():
    await conception.ecrire(ORG_CLIENT, AGENT, conception.Conception(objectif="o", reussite="r"))
    await conception.decider(ORG_CLIENT, AGENT, True, "Marie", None)


async def test_rien_n_est_rattache_sans_la_validation_d_une_personne(monde, document_cree):
    await deposer_la_faq()
    identifiant = (await executer("base_proposer", proposition(), EQUIPE))["id"]
    await approuver_la_conception()
    refus = await executer("base_rattacher", {"proposition": identifiant, "noeud": "1"}, EQUIPE)
    assert refus["ok"] is False and refus["code"] == "non_valide"
    assert (await brouillon.lire_en_attente(ORG_CLIENT, AGENT))["operations"] == [] and document_cree == []
    # Le builder n'a aucune action pour valider : ce geste est celui de la route.
    assert "base_valider" not in __import__("api.services.builder.actions", fromlist=["REGISTRE"]).REGISTRE
    reponse = client_pour(utilisateur(1)).post(f"/builder/base/{AGENT}/{identifiant}/valider")
    assert reponse.status_code == 200 and reponse.json()["document"]["document_uuid"] == "uuid-1"
    assert document_cree == [(ORG_CLIENT, "Devis et délais", proposition()["contenu"])]  # la base de connaissances du fork
    ok = await executer("base_rattacher", {"proposition": identifiant, "noeud": "1"}, EQUIPE)
    assert ok["ok"], ok
    (operation,) = (await brouillon.lire_en_attente(ORG_CLIENT, AGENT))["operations"]
    assert operation["op"] == "modifier_noeud" and operation["donnees"] == {"document_uuids": ["uuid-1"]}
    assert (await base.lister(ORG_CLIENT, AGENT))[0].rattache_a == ["1"]
    # Rattacher deux fois ne duplique pas l'identifiant.
    await executer("base_rattacher", {"proposition": identifiant, "noeud": "1"}, EQUIPE)
    derniere = (await brouillon.lire_en_attente(ORG_CLIENT, AGENT))["operations"][-1]
    assert derniere["donnees"] == {"document_uuids": ["uuid-1"]}


async def test_un_document_rejete_n_entre_pas_et_ne_se_rattache_pas(monde, document_cree):
    await deposer_la_faq()
    identifiant = (await executer("base_proposer", proposition(), EQUIPE))["id"]
    c = client_pour(utilisateur(1))
    assert c.post(f"/builder/base/{AGENT}/{identifiant}/rejeter", json={"motif": "hors sujet"}).status_code == 200
    assert c.post(f"/builder/base/{AGENT}/{identifiant}/valider").status_code == 422
    assert document_cree == []
    assert (await executer("base_rattacher", {"proposition": identifiant, "noeud": "1"}, EQUIPE))["ok"] is False


async def test_valider_revérifie_les_sources_si_le_depot_a_ete_retire(monde, document_cree):
    await deposer_la_faq()
    identifiant = (await executer("base_proposer", proposition(), EQUIPE))["id"]
    (enregistre,) = await depot.lister(ORG_CLIENT, AGENT)
    await depot.retirer(ORG_CLIENT, AGENT, enregistre["id"])
    assert client_pour(utilisateur(1)).post(f"/builder/base/{AGENT}/{identifiant}/valider").status_code == 422
    assert document_cree == []


def test_le_depot_par_les_routes_et_l_isolation_entre_organisations(monde, monkeypatch):
    c = client_pour(utilisateur(1))
    envoi = c.post(f"/builder/depot/{AGENT}", files={"fichier": ("audit.md", b"# Horaires\n\nOuvert de 9h a 18h.", "text/markdown")})
    assert envoi.status_code == 200 and envoi.json()["statut"] == "lu" and "cle_texte" not in envoi.text
    contexte = c.get(f"/builder/contexte/{AGENT}").json()
    assert [d["nom"] for d in contexte["depots"]] == ["audit.md"]
    illisible = c.post(f"/builder/depot/{AGENT}", files={"fichier": ("x.xlsx", b"PK", "application/zip")})
    assert illisible.status_code == 200 and illisible.json()["statut"] == "illisible"
    autre = client_pour(utilisateur(2, ORG_AUTRE))
    assert autre.get(f"/builder/contexte/{AGENT}").status_code == 404
    assert autre.post(f"/builder/depot/{AGENT}", files={"fichier": ("a.md", b"x", "text/plain")}).status_code == 404
    assert autre.delete(f"/builder/depot/{AGENT}/{envoi.json()['id']}").status_code == 404
    assert c.delete(f"/builder/depot/{AGENT}/{envoi.json()['id']}").status_code == 200


def test_l_adresse_d_un_site_interne_est_refusee_par_la_route(monde):
    c = client_pour(utilisateur(1))
    for adresse in ("http://169.254.169.254/latest/", "http://localhost:80/", "file:///etc/passwd"):
        reponse = c.post(f"/builder/depot/{AGENT}/site", json={"adresse": adresse})
        assert reponse.status_code == 422, adresse
    assert c.get(f"/builder/contexte/{AGENT}").json()["depots"] == []


# --- 7. la base de connaissances lit aussi PDF et Word, sur place -----------------------------------------------------------


def test_la_base_de_connaissances_lit_un_pdf_et_un_word_sans_le_service_de_l_editeur(tmp_path):
    from api.services.gen_ai.documents_locaux import est_lisible_sur_place, est_un_texte, traiter_un_document

    assert est_lisible_sur_place("a.PDF") and est_lisible_sur_place("a.docx") and not est_lisible_sur_place("a.xlsx")
    assert not est_un_texte("a.pdf", "application/pdf")  # ce n'est pas le chemin du texte : un chemin à part
    pdf = tmp_path / "audit.pdf"
    pdf.write_bytes(un_pdf([AUDIT, AUDIT_P2]))
    reponse = traiter_un_document(str(pdf), "audit.pdf", "chunked", 128)
    assert reponse["docling_metadata"]["processed_by"] == "mark-local"
    assert [c["chunk_metadata"]["page"] for c in reponse["chunks"]] == [1, 2] and "Alpha" in reponse["chunks"][1]["chunk_text"]
    mot = tmp_path / "notes.docx"
    mot.write_bytes(un_word(["Un paragraphe de l'audit."]))
    complet = traiter_un_document(str(mot), "notes.docx", "full_document", 128)
    assert "Un paragraphe" in complet["full_text"] and complet["chunks"] == []


def test_un_pdf_illisible_echoue_en_le_disant_dans_la_base_de_connaissances(tmp_path):
    from api.services.gen_ai.documents_locaux import traiter_un_document

    casse = tmp_path / "scan.pdf"
    casse.write_bytes(b"%PDF-1.4 abime")
    with pytest.raises(ValueError, match="could not be read"):
        traiter_un_document(str(casse), "scan.pdf", "chunked", 128)


# --- 8. le mode de lecture proposé (B27) et la brique de conversion unique -----------------------------------------------------


def _document_de(jetons: int) -> str:
    """Un texte qui pèse à peu près ``jetons`` jetons estimés (1,3 jeton par mot)."""
    return " ".join(["mot"] * round(jetons / 1.3))


def test_le_mode_de_lecture_se_decide_de_part_et_d_autre_du_seuil_de_4000_jetons():
    from api.services.gen_ai.documents_locaux import jetons_estimes

    assert base.SEUIL_DOCUMENT_ENTIER_JETONS == 4000
    petit, grand = _document_de(3900), _document_de(4100)
    assert jetons_estimes(petit) < 4000 <= jetons_estimes(grand)
    assert base.mode_de_lecture(petit) == "full_document"
    assert base.mode_de_lecture(grand) == "chunked"


async def test_le_mode_propose_est_celui_du_document_cree_dans_la_base_du_fork(monde, document_cree):
    texte_court = "Le devis est réalisé sur place et il est gratuit."
    await depot.deposer_fichier(ORG_CLIENT, AGENT, "faq.md", texte_court.encode(), "M")
    court = (await executer("base_proposer", {"titre": "Court", "contenu": texte_court, "sources": [{"document": "faq.md", "extrait": texte_court}]}, EQUIPE))
    assert court["ok"], court
    (document,) = await base.lister(ORG_CLIENT, AGENT)
    assert document.mode == "full_document"
    assert client_pour(utilisateur(1)).post(f"/builder/base/{AGENT}/{court['id']}/valider").status_code == 200
    assert document_cree.modes == ["full_document"]


async def test_un_long_document_part_en_recherche_par_passages(monde, document_cree):
    long = _document_de(4500)
    await depot.deposer_fichier(ORG_CLIENT, AGENT, "gros.md", long.encode(), "M")
    resultat = await executer("base_proposer", {"titre": "Long", "contenu": long, "sources": [{"document": "gros.md", "extrait": long[:200]}]}, EQUIPE)
    assert resultat["ok"], resultat
    assert client_pour(utilisateur(1)).post(f"/builder/base/{AGENT}/{resultat['id']}/valider").status_code == 200
    assert document_cree.modes == ["chunked"]


def test_une_seule_conversion_sert_au_depot_du_builder_et_a_la_base_de_connaissances(tmp_path, monkeypatch):
    """La base de connaissances n'a pas sa propre copie : elle appelle ``extraction.extraire`` (preuve : on la remplace)."""
    from api.services.builder.dossiers import extraction as module
    from api.services.gen_ai import documents_locaux

    appels = []
    vrai = module.extraire
    monkeypatch.setattr(module, "extraire", lambda nom, contenu: appels.append(nom) or vrai(nom, contenu))
    pdf = tmp_path / "audit.pdf"
    pdf.write_bytes(un_pdf([AUDIT]))
    documents_locaux.traiter_un_document(str(pdf), "audit.pdf", "chunked", 128)
    assert appels == ["audit.pdf"]


def test_la_base_de_connaissances_lit_un_texte_et_refuse_un_fichier_trop_gros(tmp_path, monkeypatch):
    from api.services.gen_ai import documents_locaux

    texte = tmp_path / "faq.txt"
    texte.write_text("Q : Ouvert ?\nR : Oui.\n", encoding="utf-8")
    assert documents_locaux.traiter_un_texte(str(texte), "chunked", 128)["chunks"]
    monkeypatch.setattr(documents_locaux, "TAILLE_MAX_OCTETS", 1000)
    gros = tmp_path / "gros.pdf"
    gros.write_bytes(un_pdf(["x" * 2000]))
    with pytest.raises(ValueError, match="too large"):
        documents_locaux.traiter_un_document(str(gros), "gros.pdf", "chunked", 128)
